import sys
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def set_style(doc):
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = 1.5
    paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1)

def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text.upper() if level == 1 else text)
    run.bold = True
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14 if level == 1 else (13 if level == 2 else 12))

def process_markdown(file_path):
    doc = Document()
    set_style(doc)
    
    with open(file_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    
    # Simple table parsing state
    in_table = False
    table = None
    
    for line in lines:
        line = line.strip()
        if not line:
            in_table = False
            continue
        
        # Handle Table
        if line.startswith('|'):
            if not in_table:
                # Start of table
                in_table = True
                cols = len([c for c in line.split('|') if c.strip()])
                table = doc.add_table(rows=0, cols=cols)
                table.style = 'Table Grid'
            
            # Add row
            if '---' in line:
                continue # Skip separator
            
            cells = [c.strip() for c in line.split('|')[1:-1]]
            row = table.add_row().cells
            for i, cell_text in enumerate(cells):
                if i < len(row):
                    row[i].text = cell_text
            continue

        in_table = False

        if line.startswith('### '):
            add_heading(doc, line[4:], 3)
        elif line.startswith('## '):
            add_heading(doc, line[3:], 2)
        elif line.startswith('# '):
            add_heading(doc, line[2:], 1)
        elif line.startswith('---'):
            doc.add_page_break()
        elif line.startswith('[IMAGE:'):
            # Parse image
            img_match = re.search(r'\[IMAGE:\s*(.*?)\]', line)
            if img_match:
                img_file = img_match.group(1)
                if os.path.exists(img_file):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = p.add_run()
                    r.add_picture(img_file, width=Inches(5.5))
                else:
                    doc.add_paragraph(f"[Missing Image: {img_file}]")
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

    doc.save('FINAL_PBL_REPORT_FINAL.docx')
    print("Successfully generated FINAL_PBL_REPORT_FINAL.docx")

if __name__ == '__main__':
    process_markdown('FINAL_REPORT.md')
